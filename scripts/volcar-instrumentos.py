# -*- coding: utf-8 -*-
"""Vuelca el contenido de todos los libros de instrumentos para revisarlos."""
import sys
from pathlib import Path

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
BASE = Path(r"F:\Maestria\tesis\taxfisco-research-lab\instrumentos")

for p in sorted(BASE.glob("*.xlsx")):
    print("#" * 100)
    print("ARCHIVO:", p.name)
    print("#" * 100)
    wb = openpyxl.load_workbook(p)
    for ws in wb:
        print(f"\n--- HOJA: {ws.title} ({ws.max_row} filas x {ws.max_column} cols) ---")
        for row in ws.iter_rows(values_only=True):
            vals = [str(v) for v in row if v is not None and str(v).strip() != ""]
            if vals:
                print(" | ".join(vals))
    print()