"""Genera un ejemplo ilustrativo de llenado del instrumento de madurez.

El documento resultante NO contiene datos de la muestra de estudio. Muestra una
sola respuesta de ejemplo, con respuestas variadas a proposito, para que quien
lo reciba vea como se marca una casilla, como se anota la justificacion y como
se calcula el nivel de un instrumento.
"""
import os
import sys

sys.stdout.reconfigure(encoding="utf-8")

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

BASE = r"F:\Maestria\tesis\taxfisco-research-lab"
SRC = os.path.join(BASE, "instrumentos",
                   "2026-06-11_Instrumento_Diagnostico_NEUTRAL_v06.docx")
if not os.path.exists(SRC):
    SRC = (r"C:\Users\fernando.mendoza\.minimax\v2\assets\2026\10\13\13-39-16-118-"
           r"asset_20261013-133916-118_14af3f8b4015_c93078e2-2026-06-11_"
           r"Instrumento_Diagnostico_NEUTRAL_v06.docx")
OUT = os.path.join(BASE, "instrumentos",
                   "2026-06-11_Instrumento_Diagnostico_EJEMPLO_LLENADO_v01.docx")

# Opcion elegida por reactivo, indice base 0. El puntaje es 4 menos el indice.
ELECCION = {
    1: 1, 2: 0, 3: 1, 4: 1,
    5: 0, 6: 1, 7: 1,
    8: 1, 9: 1, 10: 0,
    11: 2, 12: 2, 13: 2,
    14: 0, 15: 1, 16: 1,
    17: 1, 18: 1, 19: 2,
    20: 2, 21: 1,
    22: 1, 23: 1, 24: 0,
    25: 1, 26: 1, 27: 1,
    28: 1, 29: 1, 30: 2,
    31: 1, 32: 1, 33: 1,
    34: 1, 35: 1, 36: 1,
    37: 2, 38: 2, 39: 2,
    40: 1, 41: 1, 42: 1,
}

JUSTIFICACION = {
    1: "Reporte de completitud de campos del último trimestre. Muestra: 200 eventos.",
    2: "Diferencia media entre hora del evento y hora de registro, últimos 30 días.",
    3: "Percentil 95 calculado sobre la misma muestra de eventos del reactivo 1.",
    4: "Disponibilidad del servicio de monitoreo según el reporte de SLA del área.",
    5: "Muestra de 120 casos del último semestre, según registro de la mesa de ayuda.",
    6: "Porcentaje de casos con anexo probatorio verificado por el área legal.",
    7: "Conteo de tipos de evidencia detectados en el cierre de los últimos 90 días.",
    8: "Tiempo mediano de respuesta registrado en el sistema de tickets.",
    9: "Porcentaje de incidentes con respuesta dentro del tiempo objetivo.",
    10: "Número de escalamientos formalizados en el período de referencia.",
    11: "Proporción de incidentes con revisión posterior realizada y firmada.",
    12: "Número de reuniones posteriores a incidentes celebradas en el período.",
    13: "Porcentaje de cambios de proceso derivados de revisiones posteriores.",
    14: "Muestra de 90 evidencias con hash registrado en el repositorio de custodia.",
    15: "Muestra de 60 casos con cadena de custodia firmada por el custodio.",
    16: "Tiempo mediano entre la captura de la evidencia y su resguardo.",
    17: "Número de puntos de restauración ejecutados y verificados en el año.",
    18: "Porcentaje de respaldos con prueba de restauración satisfactoria.",
    19: "Número de copias de seguridad conservadas fuera de la sede principal.",
    20: "Plazo de retención efectivo según la política vigente y su aplicación.",
    21: "Solicitudes de evidencia histórica atendidas con registro de entrega.",
    22: "Porcentaje de cuentas con revisión de privilegio documentada.",
    23: "Número de cuentas privilegiadas con revisión periódica en el período.",
    24: "Porcentaje de bajas con revocación de accesos ejecutada a tiempo.",
    25: "Porcentaje de cambios validados antes de pasar a producción.",
    26: "Porcentaje de aplicaciones con revisión de seguridad realizada.",
    27: "Plazo mediano entre el reporte de vulnerabilidad y su corrección.",
    28: "Número de escaneos de seguridad ejecutados en el último mes.",
    29: "Porcentaje de activos cubiertos por el ciclo de evaluación.",
    30: "Plazo mediano de remediación de vulnerabilidades confirmadas.",
    31: "Porcentaje de eventos con los campos mínimos registrados.",
    32: "Porcentaje de registros con sello de tiempo verificable.",
    33: "Antigüedad de la evidencia en el archivo, medida por el rango de retención.",
    34: "Porcentaje de riesgos con responsable y tratamiento definido.",
    35: "Número de escenarios de riesgo actualizados en el período.",
    36: "Porcentaje de controles con riesgo residual evaluado formalmente.",
    37: "Porcentaje del personal operativo con capacitación al día.",
    38: "Número de jornadas de capacitación impartidas en el período.",
    39: "Porcentaje de cursos con evaluación de aprendizaje registrada.",
    40: "Número de indicadores en uso con reporte periódico al responsable.",
    41: "Porcentaje de indicadores con medición automatizada y sin carga manual.",
    42: "Número de decisiones de mejora sostenidas por evidencia de los indicadores.",
}

# Cortes de madurez del capitulo 4.1, expresados como techo de cada nivel.
# Inicial 0-25 %, Basico 26-50 %, Gestionado 51-75 %, Optimizado 76-100 %.
CORTES = [(25.0, "Inicial"), (50.0, "Básico"), (75.0, "Gestionado")]


def nivel(pct):
    """Devuelve el nombre del nivel de madurez que corresponde al porcentaje."""
    nombre = "Optimizado"
    for techo, etiqueta in CORTES:
        if pct <= techo:
            nombre = etiqueta
            break
    return nombre


d = Document(SRC)
body = d.element.body

# ------------------------------------------------- recorrido en orden real
tablas = []
instrumento_actual = None
orden = []
for child in body.iterchildren():
    tag = child.tag.split("}")[-1]
    if tag == "p":
        p = Paragraph(child, d)
        t = p.text.strip()
        if t.startswith("Instrumento "):
            instrumento_actual = t.split("\u2013")[0].strip()
        orden.append(("p", child, p))
    elif tag == "tbl":
        tablas.append((instrumento_actual, Table(child, d)))
        orden.append(("t", child, None))

# ------------------------------------------------- 1. marcar y justificar
for idx, (inst, t) in enumerate(tablas, start=1):
    choice = ELECCION[idx]
    for r_i, row in enumerate(t.rows):
        cells = row.cells
        if r_i == 0:
            continue
        marca = cells[1].paragraphs[0]
        marca.text = ""
        run = marca.add_run("\u2612" if r_i - 1 == choice else "\u2610")
        run.font.size = Pt(12)
        if r_i - 1 == choice:
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1F, 0x60, 0x32)
        just = cells[2].paragraphs[0]
        just.text = ""
        if r_i - 1 == choice:
            jr = just.add_run(JUSTIFICACION[idx])
            jr.font.size = Pt(9)
            jr.font.italic = True
            jr.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
        else:
            gr = just.add_run("Ejemplo: no aplica.")
            gr.font.size = Pt(8)
            gr.font.color.rgb = RGBColor(0x99, 0x9F, 0xA6)

# ------------------------------------------------- 2. puntajes por instrumento
por_inst = {}
for idx, (inst, t) in enumerate(tablas, start=1):
    por_inst.setdefault(inst, []).append(4 - ELECCION[idx])

resumen = []
for inst, pts in por_inst.items():
    tot = sum(pts)
    maximo = 4 * len(pts)
    pct = tot / maximo * 100.0
    resumen.append((inst, tot, maximo, pct, nivel(pct)))

# ------------------------------------------------- 3. rellenar lineas de puntaje
# Las lineas "Puntaje Total" aparecen en el cuerpo en el mismo orden en que se
# recorren los instrumentos, asi que se emparejan de forma secuencial.
cursor = 0
for kind, child, p in orden:
    if kind != "p":
        continue
    if not p.text.strip().startswith("Puntaje Total"):
        continue
    if cursor >= len(resumen):
        break
    inst_nombre, tot, maximo, pct, niv = resumen[cursor]
    cursor += 1
    nuevo = "Puntaje Total: %d/%d \u2192 %s%% \u2192 Nivel: %s" % (
        tot, maximo, ("%.2f" % pct).replace(".", ","), niv)
    if p.runs:
        p.runs[0].text = nuevo
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(nuevo)

# ------------------------------------------------- 4. encabezado de ejemplo
sec = d.sections[0]
hp = sec.header.paragraphs[0]
hp.text = ""
hr = hp.add_run("EJEMPLO ILUSTRATIVO DE LLENADO \u00b7 NO ES DATO DE LA MUESTRA DE ESTUDIO")
hr.font.size = Pt(8)
hr.font.bold = True
hr.font.color.rgb = RGBColor(0x9C, 0x65, 0x00)
hp.alignment = WD_ALIGN_PARAGRAPH.CENTER

fp = sec.footer.paragraphs[0]
fp.text = ""
fr = fp.add_run("Documento de ejemplo. Sirve unicamente para mostrar el procedimiento de llenado. "
                "No forma parte de la muestra de 30 respondentes ni de los resultados del estudio.")
fr.font.size = Pt(7)
fr.font.color.rgb = RGBColor(0x5A, 0x68, 0x78)
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER

# ------------------------------------------------- 5. encabezado del cuerpo
primer = d.element.body[0]
aviso = d.add_paragraph()
shd = OxmlElement("w:shd")
shd.set(qn("w:val"), "clear")
shd.set(qn("w:fill"), "FFF4E0")
aviso._p.get_or_add_pPr().append(shd)
ar = aviso.add_run("EJEMPLO ILUSTRATIVO DE LLENADO \u2014 NO ES DATO DE LA MUESTRA DE ESTUDIO\n")
ar.font.size = Pt(13)
ar.font.bold = True
ar.font.color.rgb = RGBColor(0x9C, 0x65, 0x00)
ar2 = aviso.add_run(
    "Este documento muestra cómo se marca una casilla, cómo se anota la "
    "justificación y cómo se calcula el nivel de un instrumento. Las respuestas "
    "son ilustrativas y están escolhidas a propósito con valores variados para "
    "que se observe cómo el instrumento distingue entre niveles. No proceden de "
    "respondentes reales y no deben leerse como resultado de la investigación.")
ar2.font.size = Pt(9.5)
ar2.font.color.rgb = RGBColor(0x33, 0x40, 0x4F)
primer.addprevious(aviso._p)

# ------------------------------------------------- 6. resumen del ejemplo
res = d.add_paragraph()
res.paragraph_format.space_before = Pt(14)
rr = res.add_run("Resumen del ejemplo, por instrumento")
rr.font.size = Pt(11)
rr.font.bold = True
rr.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

tabla = d.add_table(rows=1, cols=4)
tabla.style = "Table Grid"
hdr = tabla.rows[0].cells
for i, h in enumerate(["Instrumento", "Puntaje", "Cumplimiento", "Nivel"]):
    hdr[i].text = ""
    r = hdr[i].paragraphs[0].add_run(h)
    r.font.bold = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:fill"), "1F3864")
    hdr[i]._tc.get_or_add_tcPr().append(sh)

for inst_nombre, tot, maximo, pct, niv in resumen:
    cells = tabla.add_row().cells
    for i, v in enumerate([inst_nombre, "%d/%d" % (tot, maximo),
                           ("%.2f %%" % pct).replace(".", ","), niv]):
        cells[i].text = ""
        r = cells[i].paragraphs[0].add_run(v)
        r.font.size = Pt(9)

glob = sum(x[3] for x in resumen) / len(resumen)
cells = tabla.add_row().cells
for i, v in enumerate(["Madurez global del ejemplo",
                       "42 reactivos",
                       ("%.2f %%" % glob).replace(".", ","), nivel(glob)]):
    cells[i].text = ""
    r = cells[i].paragraphs[0].add_run(v)
    r.font.size = Pt(9)
    r.font.bold = True

d.save(OUT)
print("Guardado:", OUT)
print("Reactivos marcados:", len(tablas))
print("Instrumentos:", len(resumen))
print("Madurez global del ejemplo: %.2f %% -> %s" % (glob, nivel(glob)))