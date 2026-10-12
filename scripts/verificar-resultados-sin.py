# -*- coding: utf-8 -*-
"""Verificacion del libro de resultados del Prototipo II."""
import re
import openpyxl

P = (r"F:\Maestria\tesis\taxfisco-research-lab\instrumentos"
     r"\2026-10-12_Resultados_Prototipo_II_SIN_v05.xlsx")

wb = openpyxl.load_workbook(P)
print("HOJAS:", wb.sheetnames)

l = wb["Lecciones"]
abiertas = sum(1 for r in range(5, 19)
               if l.cell(row=r, column=7).value
               and "ABIERTA" in str(l.cell(row=r, column=7).value))
cerradas = 14 - abiertas
print(f"LECCIONES: {cerradas} cerradas, {abiertas} abiertas")
print("Encabezados:", [l.cell(row=4, column=c).value for c in range(1, 10)])
for r in (12, 13, 14):
    print(f"  {l.cell(row=r, column=1).value} -> "
          f"{l.cell(row=r, column=7).value} | "
          f"{str(l.cell(row=r, column=9).value)[:70]}")

res = wb["Resumen"]
for r in range(5, 26):
    a = res.cell(row=r, column=1).value
    if a and ("abiertas" in str(a) or "Ciclo" in str(a)):
        print(f"RESUMEN: {str(a)[:38]} | {res.cell(row=r, column=2).value} "
              f"{res.cell(row=r, column=3).value} | "
              f"{str(res.cell(row=r, column=4).value)[:56]}")

corrupt = []
guion = []
for ws in wb:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str):
                if re.search(r"[\u4e00-\u9fff\uf900-\ufaff\ufffd]", c.value):
                    corrupt.append((ws.title, c.coordinate))
                if c.value.strip() in ("\u2014", "-"):
                    guion.append((ws.title, c.coordinate))
print("Caracteres corruptos:", corrupt or "ninguno")
print("Celdas con guion suelto:", guion or "ninguna")

vacias = []
for name in wb.sheetnames:
    ws = wb[name]
    last = 0
    for r in range(1, ws.max_row + 1):
        if any(ws.cell(row=r, column=c).value not in (None, "")
               for c in range(1, ws.max_column + 1)):
            last = r
    for r in range(4, last + 1):
        if ws.cell(row=r, column=4).value in (None, ""):
            continue
        merged = {str(m) for m in ws.merged_cells.ranges}
        for c in range(1, ws.max_column + 1):
            if ws.cell(row=r, column=c).value in (None, ""):
                if not any(ws.cell(row=r, column=c).coordinate in m for m in merged):
                    vacias.append((name, ws.cell(row=r, column=c).coordinate))
print("Celdas vacias dentro del rango de datos:", len(vacias), vacias[:5])